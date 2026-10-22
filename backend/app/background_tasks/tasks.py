from datetime import date
import os
import calendar
import shutil

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pydantic.types import FilePath
from pydantic_settings.sources import DotEnvSettingsSource

from app.settings import settings
from app import helpers


def create_report_file(report_number: int, target_date: date):

    # Get week bounds
    week_start, week_end = helpers.get_week_bounds(target_date)

    # Build a name of the
    target_report_file = settings.BERICHTSHEFT_BASE_FILENAME.format(
        report_number=report_number,
        week_start=week_start.strftime("%d.%m.%Y"),
        week_end=week_end.strftime("%d.%m.%Y"),
    )

    # Create new report file from template and return it's path
    filepath = helpers.new_report_file_from_template(
        os.path.join(settings.REPORTS_DIRECTORY, settings.TEMPLATE_REPORT_FILENAME),
        os.path.join(settings.REPORTS_DIRECTORY, target_report_file),
    )

    # Open the report file
    report = docx.Document(filepath)

    # Set a row to work with
    header = report.sections[0].header
    table = header.tables[0]
    working_row = table.rows[1]

    # Apply new alignment for all cells in the row
    for cell in working_row.cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Fill header's cells with a sign and timestamp
    working_row.cells[0].paragraphs[0].text = settings.FULL_USER_NAME
    working_row.cells[1].paragraphs[0].text = str(report_number)
    working_row.cells[2].paragraphs[0].text = target_date.strftime("%V")
    working_row.cells[3].paragraphs[0].text = str(report_number)
    working_row.cells[4].paragraphs[0].text = week_start.strftime("%d.%m")
    working_row.cells[5].paragraphs[0].text = week_end.strftime("%d.%m")
    working_row.cells[6].paragraphs[0].text = str(date.today().year)

    # Save the report file
    report.save(filepath)


def update_report_file(target_report_filename: str, target_date: date, content: str):

    # Open and read the report file
    filepath = os.path.join(settings.REPORTS_DIRECTORY, target_report_filename)
    report = docx.Document(filepath)

    # Set a cell with a report to update
    weekday = target_date.weekday()
    working_row = 1 + weekday
    table = report.tables[0]
    row = table.rows[working_row]
    working_cell = row.cells[1]

    # Update an appropriate report
    working_cell.paragraphs[0].text = content

    # Save the report file
    report.save(filepath)
